#!/usr/bin/env python3
"""Generate updated Abhishek Purohit resume as DOCX and PDF (single page)."""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Table,
    TableStyle,
    Spacer,
    KeepTogether,
)
from reportlab.lib import colors


PROFILE = (
    "Data Analyst and Engineer with 8 years of expertise in PySpark, Python, SQL, Power BI, "
    "and Azure. Skilled in ETL pipeline design, data warehousing, and big data technologies. "
    "Adept at collaborating with cross-functional teams and communicating insights to "
    "non-technical stakeholders."
)

EXPERIENCE = [
    (
        "McKinsey &amp; Company | Data Analyst",
        "Aug 2021 – Present | Gurgaon, IN",
        [
            "Led CPG clients globally on data-driven marketing strategies, driving 16% YOY growth, $57M profitability, and price management using advanced analytics.",
            "Designed analytics for shopper behaviour; built Databricks reporting for governance, outlier detection, and Power BI KPI dashboards.",
        ],
    ),
    (
        "HCL Technologies | Engineer Programming",
        "Jun 2019 – Aug 2021 | Noida, IN",
        [
            "Automated ADF pipelines for billions of transactional rows (SQL Server, Splunk); built Power BI reports saving 12 hrs/week.",
            "Built Spark/Python ETL into Synapse/PolyBase; analysed access logs to flag 800 non-compliant users per audit standards.",
        ],
    ),
    (
        "IMSI Pvt Ltd | Associate",
        "Apr 2018 – Jun 2019 | Noida, IN",
        [
            "Developed data stack (Python, C#) on Azure/SSRS; implemented data quality checks and completed Azure Analytics POC in 3 months.",
        ],
    ),
]

TECH_STACK = (
    "<b>Data:</b> Pandas, NumPy, Alteryx, Power BI, PySpark, Databricks, Snowflake, ADF, Airflow &nbsp;|&nbsp; "
    "<b>Cloud:</b> Azure Synapse, Fabric, Functions, Logic Apps &nbsp;|&nbsp; "
    "<b>DevOps:</b> CI/CD, Docker, GitLab, Splunk &nbsp;|&nbsp; "
    "<b>DB:</b> PostgreSQL, Vertica, SQL Server, MongoDB, Delta Lake"
)

EDUCATION = [
    "<b>BITS Pilani</b> — M.Tech, Data Science &amp; Engineering (Oct 2022)",
    "<b>UTU Tehri</b> — B.Tech, Electronics &amp; Communication (July 2017)",
]

COURSES = (
    "Advanced Data Science (IBM) | Deep Learning (IIT Madras) | Machine Learning (IIT Madras) | Node.js (Pirple)"
)

CERTIFICATES = "DA 100 | DA 900 | AZ 900 | Open Hack | Applied AI"
SKILLS = "Python | JavaScript | C# | SQL"
LANGUAGES = "English | Hindi"


def build_docx(path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    def set_spacing(paragraph, before=0, after=2, line=11):
        pf = paragraph.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = Pt(line)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Abhishek Purohit")
    run.bold = True
    run.font.size = Pt(14)
    set_spacing(title, after=0)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Data Engineer & Analyst | +91-7060470209 | abhishek.purohit7ap@gmail.com"
    )
    run.font.size = Pt(8.5)
    set_spacing(subtitle, after=4)

    def section_heading(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        set_spacing(p, before=3, after=1, line=10)
        return p

    def body(text, size=8, bold=False, italic=False):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        set_spacing(p, after=1, line=10)
        return p

    def bullet(text):
        p = doc.add_paragraph(text, style="List Bullet")
        set_spacing(p, after=0, line=10)
        for run in p.runs:
            run.font.size = Pt(8)
        return p

    section_heading("PROFILE")
    body(PROFILE, size=8)

    section_heading("EXPERIENCE")
    for company, dates, bullets in EXPERIENCE:
        body(company.replace("&amp;", "&"), size=8, bold=True)
        body(dates.replace("&amp;", "&"), size=8, italic=True)
        for item in bullets:
            bullet(item)

    section_heading("TECH STACK")
    body(TECH_STACK.replace("<b>", "").replace("</b>", "").replace("&nbsp;|&nbsp;", " | "), size=8)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    left, right = table.rows[0].cells

    left.paragraphs[0].add_run("EDUCATION").bold = True
    for item in EDUCATION:
        p = left.add_paragraph(item.replace("<b>", "").replace("</b>", "").replace("&amp;", "&"))
        for run in p.runs:
            run.font.size = Pt(8)
        set_spacing(p, after=0, line=10)

    right.paragraphs[0].add_run("CERTIFICATES & SKILLS").bold = True
    for line in [CERTIFICATES, SKILLS, LANGUAGES]:
        p = right.add_paragraph(line)
        for run in p.runs:
            run.font.size = Pt(8)
        set_spacing(p, after=0, line=10)

    section_heading("COURSES")
    body(COURSES, size=8)

    doc.save(path)


def build_pdf(path):
    styles = getSampleStyleSheet()
    page_w, page_h = letter
    margin = 0.42 * inch
    content_w = page_w - 2 * margin

    title_style = ParagraphStyle(
        "Title", parent=styles["Normal"], fontSize=13, alignment=1, spaceAfter=1, leading=14
    )
    subtitle_style = ParagraphStyle(
        "Subtitle", parent=styles["Normal"], fontSize=8, alignment=1, spaceAfter=5, leading=9
    )
    section_style = ParagraphStyle(
        "Section",
        parent=styles["Normal"],
        fontSize=8.5,
        fontName="Helvetica-Bold",
        spaceBefore=3,
        spaceAfter=1,
        leading=10,
    )
    body_style = ParagraphStyle(
        "Body", parent=styles["Normal"], fontSize=7.5, spaceAfter=1, leading=9
    )
    job_title_style = ParagraphStyle(
        "JobTitle",
        parent=styles["Normal"],
        fontSize=7.5,
        fontName="Helvetica-Bold",
        spaceBefore=2,
        spaceAfter=0,
        leading=9,
    )
    job_date_style = ParagraphStyle(
        "JobDate",
        parent=styles["Normal"],
        fontSize=7.5,
        fontName="Helvetica-Oblique",
        spaceAfter=0,
        leading=9,
    )
    bullet_style = ParagraphStyle(
        "Bullet",
        parent=styles["Normal"],
        fontSize=7.5,
        leftIndent=10,
        bulletIndent=4,
        spaceAfter=0,
        leading=9,
    )

    story = [
        Paragraph("Abhishek Purohit", title_style),
        Paragraph(
            "Data Engineer &amp; Analyst | +91-7060470209 | abhishek.purohit7ap@gmail.com",
            subtitle_style,
        ),
        Paragraph("PROFILE", section_style),
        Paragraph(PROFILE, body_style),
        Paragraph("EXPERIENCE", section_style),
    ]

    for company, dates, bullets in EXPERIENCE:
        block = [
            Paragraph(company, job_title_style),
            Paragraph(dates, job_date_style),
        ]
        for item in bullets:
            block.append(Paragraph(f"• {item}", bullet_style))
        story.append(KeepTogether(block))

    story.append(Paragraph("TECH STACK", section_style))
    story.append(Paragraph(TECH_STACK, body_style))

    col_w = content_w / 2 - 6
    left_col = [
        Paragraph("EDUCATION", section_style),
    ]
    for item in EDUCATION:
        left_col.append(Paragraph(item, body_style))

    right_col = [
        Paragraph("CERTIFICATES &amp; SKILLS", section_style),
        Paragraph(CERTIFICATES, body_style),
        Paragraph(SKILLS, body_style),
        Paragraph(f"<b>Languages:</b> {LANGUAGES}", body_style),
    ]

    two_col = Table([[left_col, right_col]], colWidths=[col_w, col_w])
    two_col.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    story.append(two_col)

    story.append(Paragraph("COURSES", section_style))
    story.append(Paragraph(COURSES, body_style))

    doc = SimpleDocTemplate(
        path,
        pagesize=letter,
        rightMargin=margin,
        leftMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
    )
    doc.build(story)


def page_count(pdf_path):
    try:
        from pypdf import PdfReader

        return len(PdfReader(pdf_path).pages)
    except ImportError:
        import subprocess

        result = subprocess.run(
            ["python3", "-c", f"from pypdf import PdfReader; print(len(PdfReader('{pdf_path}').pages))"],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0:
            pip_install = subprocess.run(["pip", "install", "pypdf", "-q"], capture_output=True)
            if pip_install.returncode == 0:
                from pypdf import PdfReader

                return len(PdfReader(pdf_path).pages)
        return None


if __name__ == "__main__":
    pdf_path = "/workspace/resume/Abhishek_Purohit_Data_Engineer.pdf"
    docx_path = "/workspace/resume/Abhishek_Purohit_Data_Engineer.docx"
    build_docx(docx_path)
    build_pdf(pdf_path)
    pages = page_count(pdf_path)
    print(f"Resume generated successfully. PDF pages: {pages}")
